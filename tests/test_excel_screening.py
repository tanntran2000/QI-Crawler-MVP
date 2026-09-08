from __future__ import annotations

from dataclasses import replace
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from qi_crawler.gui_services import run_bid_radar_excel_screening
from qi_crawler.market_intelligence import excel_screening as screening
from qi_crawler.market_intelligence.excel_screening import (
    C07Result,
    C09Result,
    C10Result,
    OperationalCriterionResult,
    OperationalResult,
    ScreeningResult,
    export_screening_workbook,
    screen_excel_workbook,
)


def _write_khmt(path, rows, *, filename_header="TBMT_07_09.xlsx"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "KHLCNT"
    sheet.append(["Báo cáo kế hoạch lựa chọn nhà thầu"])
    sheet.append([])
    sheet.append(
        [
            "SỐ KẾ HOẠCH",
            "TÊN GÓI THẦU",
            "TÊN DỰ ÁN",
            "TÊN CHỦ ĐẦU TƯ",
            "GIÁ GÓI THẦU",
            "NỘI DUNG PHÊ DUYỆT",
        ]
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_tbmt(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bản tin điện tử"
    sheet.append(["TBMT report"])
    sheet.append([])
    sheet.append(["BÊN MỜI THẦU", "DỰ ÁN", "GÓI THẦU", "GIÁ GÓI THẦU", "ĐỊA CHỈ BÊN MỜI THẦU"])
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _write_operational_tbmt(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bản tin điện tử"
    sheet.append(["TBMT report"])
    sheet.append([])
    sheet.append(
        [
            "BÊN MỜI THẦU",
            "DỰ ÁN",
            "GÓI THẦU",
            "GIÁ GÓI THẦU",
            "ĐỊA CHỈ BÊN MỜI THẦU",
            "HÌNH THỨC LỰA CHỌN NHÀ THẦU",
            "PHƯƠNG THỨC LỰA CHỌN NHÀ THẦU",
            "ĐỊA ĐIỂM THỰC HIỆN",
        ]
    )
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def _khmt_row(plan_id, package, project, buyer, price, approval=""):
    return [plan_id, package, project, buyer, price, approval]


def test_source_authority_uses_identity_and_accounts_duplicate_packages(tmp_path):
    path = tmp_path / "TBMT_07_09.xlsx"
    _write_khmt(
        path,
        [
            _khmt_row(
                "PL2600290673-00",
                "Mua máy chủ và thiết bị lưu trữ",
                "Chương trình chuyển đổi số",
                "Đơn vị Đồng Nai",
                "1.999.999.999",
            ),
            _khmt_row(
                "PL2600290673-00",
                "Mua sắm trang thiết bị",
                "Chương trình chuyển đổi số",
                "Đơn vị Đồng Nai",
                "2.000.000.001",
            ),
        ],
    )
    result = screen_excel_workbook(path)

    assert result.source_type == "KHMT"
    assert result.filename_schema_conflict is True
    assert result.identity_namespace == "PL"
    assert result.data_record_rows == 2
    assert result.read_ok_rows == 2
    assert len(result.records) == 2
    assert [record.identity for record in result.records] == [
        "PL2600290673-00",
        "PL2600290673-00",
    ]
    assert result.records[0].c07 is C07Result.PASS
    assert result.records[0].c09 is C09Result.PASS
    assert result.records[1].c07 is C07Result.UNKNOWN
    assert result.records[1].screening is ScreeningResult.NEEDS_REVIEW
    assert result.records[0].c10 is C10Result.UNKNOWN
    assert result.records[0].location_hint == "Đồng Nai"


def test_c07_is_package_scoped_and_generic_equipment_stays_unknown(tmp_path):
    path = tmp_path / "KHMT.xlsx"
    _write_khmt(
        path,
        [
            _khmt_row("PL2600000001-00", "Triển khai mạng LAN và switch", "", "", 10),
            _khmt_row("PL2600000002-00", "Mua sắm trang thiết bị", "Chương trình chuyển đổi số", "", 10),
            _khmt_row("PL2600000003-00", "Cấp license Windows", "", "", 10),
            _khmt_row("PL2600000004-00", "Thiết bị an toàn thông tin firewall", "", "", 10),
            _khmt_row("PL2600000005-00", "Mua camera giám sát", "", "", 10),
        ],
    )
    result = screen_excel_workbook(path)

    assert [record.c07 for record in result.records] == [
        C07Result.PASS,
        C07Result.UNKNOWN,
        C07Result.PASS,
        C07Result.PASS,
        C07Result.PASS,
    ]
    assert result.records[1].screening is ScreeningResult.NEEDS_REVIEW
    assert result.records[0].c07_rule_code == "R07-NETWORK"
    assert result.records[2].c07_rule_code == "R07-SOFTWARE-LICENSE"
    assert result.records[3].c07_rule_code == "R07-CYBERSECURITY"


def test_c09_exact_two_billion_boundary(tmp_path):
    path = tmp_path / "KHMT.xlsx"
    _write_khmt(
        path,
        [
            _khmt_row("PL2600000011-00", "Mua máy tính", "", "", "1.999.999.999"),
            _khmt_row("PL2600000012-00", "Mua máy tính", "", "", "2.000.000.000"),
            _khmt_row("PL2600000013-00", "Mua máy tính", "", "", "2.000.000.001"),
            _khmt_row("PL2600000014-00", "Mua máy tính", "", "", "không rõ"),
        ],
    )
    result = screen_excel_workbook(path)

    assert [record.package_price for record in result.records] == [
        Decimal(1999999999),
        Decimal(2000000000),
        Decimal(2000000001),
        None,
    ]
    assert [record.c09 for record in result.records] == [
        C09Result.PASS,
        C09Result.PASS,
        C09Result.FAIL,
        C09Result.UNKNOWN,
    ]
    assert all(record.screening is ScreeningResult.SELECT for record in result.records)


def test_accounting_and_four_sheet_export_preserve_provenance(tmp_path):
    source = tmp_path / "KHMT.xlsx"
    _write_khmt(
        source,
        [
            _khmt_row("PL2600000021-00", "Mua máy tính", "", "", 100),
            ["không phải mã PL", "Gói bị lỗi", "", "", 100, ""],
        ],
    )
    workbook = load_workbook(source)
    workbook.active.append(["TỔNG CỘNG", None, None, None, None, None])
    workbook.save(source)

    result = screen_excel_workbook(source)
    assert result.post_header_nonempty_rows == 3
    assert result.data_record_rows == 2
    assert result.read_ok_rows == 1
    assert result.read_error_rows == 1
    assert result.non_record_rows == 1
    assert result.accounting_invariant_status == "PASS"

    output = tmp_path / "screened.xlsx"
    export_screening_workbook(result, output)
    exported = load_workbook(output, read_only=True, data_only=True)
    assert exported.sheetnames == [
        "01_CO_HOI_XEM_XET",
        "02_CAN_BO_SUNG",
        "03_AUDIT_TOAN_BO",
        "04_META",
    ]
    assert next(exported["01_CO_HOI_XEM_XET"].iter_rows(values_only=True))[:6] == (
        "Tên gói",
        "Giá gói",
        "Lý do phù hợp",
        "Địa bàn gợi ý",
        "Cần kiểm tra gì",
        "Mã PL/IB",
    )
    audit_values = list(exported["03_AUDIT_TOAN_BO"].iter_rows(values_only=True))
    assert any(row[0] == "READ_ERROR" for row in audit_values[1:])
    meta = {row[0]: row[1] for row in exported["04_META"].iter_rows(values_only=True) if row[0]}
    assert meta["SOURCE_SHA256"] == result.source_sha256
    assert meta["DATA_RECORD_ROWS"] == 2
    assert meta["READ_ERROR_ROWS"] == 1
    exported.close()


def test_tbmt_identity_schema_is_retained_and_location_is_hint_only(tmp_path):
    path = tmp_path / "KHMT_named_wrongly.xlsx"
    _write_tbmt(
        path,
        [
            [
                "Bên mời thầu",
                "Dự án số hóa",
                "Mua phần mềm bản quyền (Số thông báo: IB2600500001-00)",
                "2.000.000.000",
                "TP.HCM",
            ]
        ],
    )
    result = screen_excel_workbook(path)
    record = result.records[0]
    assert result.source_type == "TBMT"
    assert result.filename_schema_conflict is True
    assert result.identity_namespace == "IB"
    assert record.identity == "IB2600500001-00"
    assert record.c07 is C07Result.PASS
    assert record.c10 is C10Result.UNKNOWN
    assert record.location_hint == "TPHCM"


def test_gui_service_delegates_screening_and_export(tmp_path):
    source = tmp_path / "KHMT.xlsx"
    _write_khmt(source, [_khmt_row("PL2600000031-00", "Mua máy tính", "", "", 100)])
    output = tmp_path / "team-bid-screening.xlsx"
    result = run_bid_radar_excel_screening(source, output_path=output)

    assert result.output_path == output.resolve()
    assert result.run.select_count == 1
    assert output.is_file()


def test_operational_profile_evaluates_hard_criteria_without_changing_canonical_results(tmp_path):
    path = tmp_path / "TBMT_operational.xlsx"
    _write_operational_tbmt(
        path,
        [
            [
                "Buyer",
                "Project",
                "Cung cấp máy vi tính và OCR (Số thông báo: IB2600000101-00)",
                "1.000.000.000",
                "Hà Nội",
                "Đấu thầu rộng rãi trong nước",
                "Một giai đoạn một túi hồ sơ",
                "TPHCM",
            ],
            [
                "Buyer",
                "Chương trình chuyển đổi số",
                "Mua sắm trang thiết bị (Số thông báo: IB2600000102-00)",
                "1.000.000.000",
                "TPHCM",
                "Đấu thầu rộng rãi trong nước",
                "Một giai đoạn một túi hồ sơ",
                "TPHCM",
            ],
            [
                "Buyer",
                "Project",
                "Cung cấp máy tính xách tay (Số thông báo: IB2600000103-00)",
                "3.000.000.000",
                "TPHCM",
                "Đấu thầu rộng rãi trong nước",
                "Một giai đoạn một túi hồ sơ",
                "TPHCM",
            ],
            [
                "Buyer",
                "Project",
                "Cung cấp máy vi tính (Số thông báo: IB2600000104-00)",
                "1.000.000.000",
                "TPHCM",
                "Chỉ định thầu",
                "Một giai đoạn một túi hồ sơ",
                "TPHCM",
            ],
            [
                "Buyer",
                "Project",
                "Cung cấp máy vi tính (Số thông báo: IB2600000105-00)",
                "1.000.000.000",
                "TPHCM",
                "Đấu thầu rộng rãi trong nước",
                "",
                "TPHCM",
            ],
            [
                "Buyer",
                "Project",
                "Cung cấp máy vi tính (Số thông báo: IB2600000106-00)",
                "1.000.000.000",
                "TPHCM",
                "Đấu thầu rộng rãi trong nước",
                "Một giai đoạn một túi hồ sơ",
                "",
            ],
        ],
    )

    result = screen_excel_workbook(path, profile=replace(screening.QI_TEAM_BID_PROFILE_V1, allowed_execution_locations=("TPHCM",)))

    assert result.records[0].c07 is C07Result.PASS
    assert result.records[0].screening is ScreeningResult.SELECT
    assert result.records[0].operational_result is OperationalResult.FIT
    assert result.records[0].location_status is OperationalCriterionResult.PASS
    assert result.records[1].c07 is C07Result.UNKNOWN
    assert result.records[1].operational_result is OperationalResult.NEEDS_INFO
    assert result.records[2].c09 is C09Result.FAIL
    assert result.records[2].budget_priority is screening.BudgetPriority.EXPANSION_OPPORTUNITY
    assert result.records[2].operational_result is OperationalResult.FIT
    assert result.records[3].selection_status is OperationalCriterionResult.FAIL
    assert result.records[3].operational_result is OperationalResult.OUTSIDE_PROFILE
    assert result.records[4].procurement_status is OperationalCriterionResult.UNKNOWN
    assert result.records[4].operational_result is OperationalResult.NEEDS_INFO
    assert result.records[5].location_status is OperationalCriterionResult.UNKNOWN
    assert result.records[5].operational_result is OperationalResult.NEEDS_INFO


def test_c07_extension_is_bounded_and_contextual(tmp_path):
    path = tmp_path / "KHMT_c07_extension.xlsx"
    _write_khmt(
        path,
        [
            _khmt_row("PL2600000101-00", "Cung cấp NAS và tủ rack", "", "", 10),
            _khmt_row("PL2600000102-00", "Cung cấp DDoS và WAN", "", "", 10),
            _khmt_row("PL2600000103-00", "Số hóa tài liệu", "", "", 10),
            _khmt_row("PL2600000104-00", "Số hóa tài liệu bằng OCR và cơ sở dữ liệu", "", "", 10),
            _khmt_row("PL2600000105-00", "Mua sắm trang thiết bị", "Mô hình chuyển đổi số", "", 10),
        ],
    )
    result = screen_excel_workbook(path)

    assert result.records[0].c07 is C07Result.PASS
    assert result.records[1].c07 is C07Result.PASS
    assert result.records[2].c07 is C07Result.UNKNOWN
    assert result.records[3].c07 is C07Result.PASS
    assert result.records[4].c07 is C07Result.UNKNOWN


def test_operational_export_partitions_profile_results_and_meta(tmp_path):
    source = tmp_path / "TBMT_operational.xlsx"
    _write_operational_tbmt(
        source,
        [
            ["Buyer", "Project", "Cung cấp máy vi tính (Số thông báo: IB2600000201-00)", "100", "", "Đấu thầu rộng rãi", "Một giai đoạn một túi hồ sơ", "TPHCM"],
            ["Buyer", "Project", "Mua sắm trang thiết bị (Số thông báo: IB2600000202-00)", "100", "", "Đấu thầu rộng rãi", "Một giai đoạn một túi hồ sơ", "TPHCM"],
            ["Buyer", "Project", "Cung cấp máy vi tính (Số thông báo: IB2600000203-00)", "3.000.000.000", "", "Đấu thầu rộng rãi", "Một giai đoạn một túi hồ sơ", "TPHCM"],
        ],
    )
    run = screen_excel_workbook(source, profile=replace(screening.QI_TEAM_BID_PROFILE_V1, allowed_execution_locations=("TPHCM",)))
    output = tmp_path / "operational.xlsx"
    export_screening_workbook(run, output)
    exported = load_workbook(output, read_only=True, data_only=True)

    fit_rows = list(exported["01_CO_HOI_XEM_XET"].iter_rows(values_only=True))
    review_rows = list(exported["02_CAN_BO_SUNG"].iter_rows(values_only=True))
    audit_rows = list(exported["03_AUDIT_TOAN_BO"].iter_rows(values_only=True))
    meta = {row[0]: row[1] for row in exported["04_META"].iter_rows(values_only=True) if row[0]}

    assert len(fit_rows) == 3
    assert len(review_rows) == 2
    assert len(audit_rows) == 4
    assert meta["OPERATIONAL_PROFILE_ID"] == "QI_TEAM_BID_PROFILE_V1"
    assert meta["FIT_COUNT"] == 2
    assert meta["NEEDS_INFO_COUNT"] == 1
    assert meta["OUTSIDE_PROFILE_COUNT"] == 0
    exported.close()


@pytest.mark.parametrize("token", ["IB2600486943- 01", "ib2600486943 - 01", "IB2600486943\t-\t01"])
def test_screening_canonical_identity_keeps_raw_source(tmp_path, token):
    path = tmp_path / "TBMT.xlsx"
    title = f"Mua máy tính (Số thông báo: {token})"
    _write_tbmt(path, [["Buyer", "", title, 100, ""]])
    record = screen_excel_workbook(path).records[0]
    assert record.identity == "IB2600486943-01"
    assert screening._revision(record.identity) == "01"
    assert record.raw_fields["GÓI THẦU"] == title


def test_geography_needs_explicit_profile_and_preserves_evidence():
    fields = {"ĐỊA ĐIỂM THỰC HIỆN": "TPHCM"}
    assert screening._screen_execution_location(fields)[0] == "UNKNOWN"
    configured = replace(screening.QI_TEAM_BID_PROFILE_V1, allowed_execution_locations=("TPHCM",))
    assert screening._screen_execution_location(fields, configured) == ("PASS", "TPHCM")
    assert screening._screen_execution_location({"ĐỊA CHỈ BÊN MỜI THẦU": "TPHCM"}, configured)[0] == "UNKNOWN"
    assert screening._screen_execution_location({"ĐỊA ĐIỂM THỰC HIỆN": "Hà Nội"}, configured)[0] == "FAIL"
    assert screening._screen_execution_location({"ĐỊA ĐIỂM THỰC HIỆN": "Đường Hồ Chí Minh"}, configured)[0] == "UNKNOWN"


@pytest.mark.parametrize("title, expected", [
    ("Mua linh kiện điện tử", "PASS"), ("Nâng cấp hạ tầng công nghệ thông tin", "PASS"),
    ("Nâng cấp hạ tầng CNTT", "PASS"), ("Mua trang thiết bị công nghệ thông tin", "PASS"),
    ("Mua thiết bị CNTT", "PASS"), ("Cung cấp Mini rack tại IDC HLC", "PASS"),
    ("Mua rack trưng bày hàng hóa", "UNKNOWN"), ("Mini rack", "UNKNOWN"),
    ("Số hóa tài liệu", "UNKNOWN"), ("mạng công nghệ mô hình lắp đặt", "UNKNOWN"),
    ("Dịch vụ đào tạo Ứng dụng phần mềm mô phỏng", "UNKNOWN"),
    ("Tập huấn sử dụng thiết bị CNTT", "UNKNOWN"),
    ("Training Office", "UNKNOWN"),
    ("Đào tạo sử dụng phần mềm; cung cấp bản quyền phần mềm", "PASS"),
    ("Cung cấp phần mềm và đào tạo sử dụng", "PASS"),
])
def test_forward_c07_deliverable_precision(title, expected):
    assert screening._screen_c07(screening.SourceType.TBMT, {"GÓI THẦU": title})[0] == expected


def test_training_title_cannot_leak_via_content_or_project():
    fields = {"GÓI THẦU": "Đào tạo sử dụng phần mềm", "NỘI DUNG CHÍNH CỦA GÓI THẦU": "Nội dung phần mềm Office", "DỰ ÁN": "Mua máy chủ"}
    assert screening._screen_c07(screening.SourceType.TBMT, fields)[0] == "UNKNOWN"


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Đầu tư, nâng cấp hạ tầng công nghệ thông tin tại Sở Giáo dục và Đào tạo", "PASS"),
        ("Thuê trang thiết bị và dịch vụ CNTT cho công trình Trường Đại học Công nghệ", "PASS"),
        ("Cung cấp phần mềm quản lý và đào tạo chuyển giao sử dụng", "PASS"),
        ("Dịch vụ đào tạo Ứng dụng phần mềm mô phỏng", "UNKNOWN"),
    ],
)
def test_c07_distinguishes_package_deliverable_from_organization_or_training(title, expected):
    assert screening._screen_c07(screening.SourceType.TBMT, {"GÓI THẦU": title})[0] == expected


@pytest.mark.parametrize(
    ("price", "expected"),
    [
        ("1.999.999.999", "CURRENT_PRIORITY"),
        ("2.000.000.000", "CURRENT_PRIORITY"),
        ("2.000.000.001", "EXPANSION_OPPORTUNITY"),
        (None, "UNKNOWN"),
    ],
)
def test_budget_priority_is_configurable_without_becoming_a_hard_exclusion(tmp_path, price, expected):
    source = tmp_path / "TBMT_budget_priority.xlsx"
    _write_operational_tbmt(source, [[
        "Buyer", "Project", "Cung cấp máy vi tính (Số thông báo: IB2600000301-00)", price,
        "", "Đấu thầu rộng rãi", "Một giai đoạn một túi hồ sơ", "TPHCM",
    ]])
    profile = replace(screening.QI_TEAM_BID_PROFILE_V1, allowed_execution_locations=("TPHCM",))
    record = screen_excel_workbook(source, profile=profile).records[0]

    assert record.budget_priority.value == expected
    assert record.operational_result is OperationalResult.FIT


def test_business_export_exposes_suitability_and_budget_priority_separately(tmp_path):
    source = tmp_path / "TBMT_budget_priority_export.xlsx"
    _write_operational_tbmt(source, [[
        "Buyer", "Project", "Cung cấp máy vi tính (Số thông báo: IB2600000401-00)", "2.000.000.001",
        "", "Đấu thầu rộng rãi", "Một giai đoạn một túi hồ sơ", "TPHCM",
    ]])
    profile = replace(screening.QI_TEAM_BID_PROFILE_V1, allowed_execution_locations=("TPHCM",))
    output = tmp_path / "budget-priority.xlsx"
    export_screening_workbook(screen_excel_workbook(source, profile=profile), output)
    book = load_workbook(output, read_only=True, data_only=True)

    headers = next(book["01_CO_HOI_XEM_XET"].values)
    row = next(book["01_CO_HOI_XEM_XET"].iter_rows(min_row=2, values_only=True))
    assert "Mức phù hợp của gói thầu" in headers
    assert "Ưu tiên ngân sách Team Bid" in headers
    assert row[headers.index("Mức phù hợp của gói thầu")] == "Phù hợp hồ sơ sàng lọc"
    assert row[headers.index("Ưu tiên ngân sách Team Bid")] == "Cơ hội mở rộng"
    assert book["02_CAN_BO_SUNG"].max_row == 1
    book.close()


def test_review_sheet_only_needs_info_sorted_by_missing_location(tmp_path):
    p = tmp_path / "TBMT.xlsx"
    _write_operational_tbmt(p, [
        ["B", "", "Mua sắm (Số thông báo: IB2600000001-00)", 100, "", "", "", ""],
        ["B", "", "Mua máy tính (Số thông báo: IB2600000002-00)", None, "", "", "", ""],
        ["B", "", "Mua sắm (Số thông báo: IB2600000003-00)", 3000000000, "", "", "", ""],
        ["B", "", "Mua máy tính (Số thông báo: IB2600000004-00)", 100, "", "Đấu thầu rộng rãi", "Một giai đoạn một túi hồ sơ", ""],
    ])
    run = screen_excel_workbook(p)
    records = screening._ordered_operational_review(run.records)
    assert [r.source_row for r in records] == [7, 5, 4, 6]
    out = tmp_path / "out.xlsx"
    export_screening_workbook(run, out)
    book = load_workbook(out, read_only=True, data_only=True)
    headers = next(book["02_CAN_BO_SUNG"].values)
    assert "C09" not in headers
    assert "Ưu tiên ngân sách Team Bid" in headers
    assert book["02_CAN_BO_SUNG"].max_row == 5
    assert book["03_AUDIT_TOAN_BO"].max_row == 5
    assert "Cần bổ sung thông tin" in next(book["02_CAN_BO_SUNG"].iter_rows(min_row=2, values_only=True))
    book.close()
