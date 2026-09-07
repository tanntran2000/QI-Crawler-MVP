from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook


def _builder_module():
    path = Path(__file__).parents[1] / "tools" / "ground_truth" / "build_golden_candidate_01.py"
    spec = importlib.util.spec_from_file_location("golden_candidate_builder", path)
    if spec is None or spec.loader is None:
        pytest.fail("ground-truth builder module is not available")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


HEADERS = (
    "GÓI TIN",
    "BÊN MỜI THẦU",
    "ĐỊA CHỈ BÊN MỜI THẦU",
    "DỰ ÁN",
    "GÓI THẦU",
    "NỘI DUNG CHÍNH CỦA GÓI THẦU",
    "NGUỒN VỐN",
    "GIÁ GÓI THẦU",
    "PHƯƠNG THỨC LỰA CHỌN NHÀ THẦU",
    "HÌNH THỨC LỰA CHỌN NHÀ THẦU",
    "THỜI GIAN PHÁT HÀNH HSMT",
    "GIÁ BÁN 1 BỘ HSMT",
    "BẢO ĐẢM DỰ THẦU",
    "HÌNH THỨC BẢO ĐẢM DỰ THẦU",
    "ĐỊA ĐIỂM PHÁT HÀNH",
    "THỜI GIAN ĐÓNG THẦU(HẠN CUỐI TIẾP NHẬN BG)",
    "THỜI GIAN MỞ THẦU",
    "THỜI GIAN THỰC HIỆN HỢP ĐỒNG",
)

IDENTITIES = (
    "IB2600493305-01",
    "IB2600508933-00",
    "IB2600510657-01",
    "IB2600510453-00",
    "IB2600510121-00",
    "IB2600510139-00",
    "IB2600510867-00",
    "IB2600506108-00",
    "IB2600510159-00",
    "IB2600508572-00",
    "IB2600510970-00",
)

EXPECTED_HUMAN_REGIONS = (
    "TP.HCM",
    "Cần Giờ",
    "Vũng Tàu",
    "Bình Dương",
    "Tây Ninh",
    "Long An",
    "Đồng Nai",
    "Lâm Đồng",
    "Phan Thiết",
)

UNAUTHORIZED_REGIONS = (
    "Hà Nội",
    "Hải Phòng",
    "Quảng Ninh",
    "Bắc Ninh",
    "Hưng Yên",
    "Hải Dương",
    "Vĩnh Phúc",
    "Hà Nam",
    "Nam Định",
)


def _write_source(path: Path, *, identities: tuple[str, ...] = IDENTITIES) -> str:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bản tin điện tử"
    for _ in range(9):
        sheet.append([None] * len(HEADERS))
    sheet.append(list(HEADERS))
    for index, raw_id in enumerate(identities, start=1):
        values = [None] * len(HEADERS)
        values[0] = "1. Thông báo mời thầu"
        values[1] = "Bên mời thầu"
        values[2] = "Địa chỉ bên mời thầu"
        values[3] = "Mua sắm thiết bị mạng"
        values[4] = f"Mua sắm thiết bị mạng (Số thông báo: {raw_id})"
        values[5] = ""
        values[6] = "Ngân sách Nhà nước"
        values[7] = "1.000.000.000"
        values[8] = "Một giai đoạn một túi hồ sơ"
        values[9] = "Đấu thầu rộng rãi trong nước"
        values[10] = "16 giờ 11 ngày 05/09/2026"
        values[12] = "50.000.000"
        values[13] = "Thư bảo lãnh"
        values[14] = "https://muasamcong.mpi.gov.vn"
        values[15] = "07 giờ 00 ngày 16/09/2026"
        values[16] = "07 giờ 00 ngày 16/09/2026"
        values[17] = "90 ngày"
        sheet.append(values)
    workbook.save(path)
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_builder_generates_four_sheet_draft_with_source_and_labeling_rows(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    output = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT.xlsx"
    source_sha = _write_source(source)

    result = module.build_draft(source, output, expected_source_sha256=source_sha)

    assert result.source_data_rows == 11
    assert result.successfully_imported_rows == 11
    assert result.rejected_or_skipped_rows == 0
    assert result.import_issue_count == 0
    assert result.exact_identity_count == 11
    assert result.output == output.resolve()
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == [
        "00_HUONG_DAN_PROFILE",
        "01_NGUON_11_GOI",
        "02_TEAM_BID_GAN_NHAN",
        "03_BANG_CHUNG_TIEU_CHI",
    ]
    assert workbook["01_NGUON_11_GOI"].max_row == 12
    assert workbook["02_TEAM_BID_GAN_NHAN"].max_row == 12
    assert workbook["03_BANG_CHUNG_TIEU_CHI"].max_row == 122
    workbook.close()


def test_builder_does_not_prefill_business_answers_or_predictions(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    output = tmp_path / "draft.xlsx"
    source_sha = _write_source(source)

    module.build_draft(source, output, expected_source_sha256=source_sha)

    workbook = load_workbook(output, data_only=False)
    profile = workbook["00_HUONG_DAN_PROFILE"]
    assert profile["B2"].value == "GOLDEN_CANDIDATE_01"
    assert profile["B3"].value == "DRAFT"
    assert profile["B10"].value in (None, "")
    assert profile["B11"].value in (None, "")
    assert profile["B12"].value == "Asia/Ho_Chi_Minh"
    assert profile["B13"].value in (None, "")
    labels = workbook["02_TEAM_BID_GAN_NHAN"]
    headers = {cell.value: cell.column for cell in labels[1]}
    assert all(labels.cell(row, headers["EXPECTED_SCREENING"]).value in (None, "") for row in range(2, 13))
    assert all(labels.cell(row, headers["EXPECTED_PRIORITY"]).value in (None, "") for row in range(2, 13))
    evidence = workbook["03_BANG_CHUNG_TIEU_CHI"]
    evidence_headers = {cell.value: cell.column for cell in evidence[1]}
    assert all(evidence.cell(row, evidence_headers["EXPECTED_OUTCOME"]).value in (None, "") for row in range(2, 123))
    assert "CRAWLER_RESULT" not in {cell.value for cell in labels[1]}
    workbook.close()


def test_builder_preserves_provenance_and_known_anchor(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    output = tmp_path / "draft.xlsx"
    source_sha = _write_source(source)

    module.build_draft(source, output, expected_source_sha256=source_sha)

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["01_NGUON_11_GOI"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(4, headers["QI_IB_RAW"]).value == "IB2600510657-01"
    assert sheet.cell(4, headers["QI_SOURCE_ROW"]).value == 13
    assert sheet.cell(4, headers["QI_SOURCE_SHA256"]).value == source_sha
    assert sheet.cell(4, headers["QI_OBSERVATION_KEY"]).value
    assert sheet.cell(4, headers["QI_SELECTION_METHOD_NORMALIZED"]).value == "DAU_THAU_RONG_RAI"
    workbook.close()


def test_builder_rejects_source_identity_and_output_safety_violations(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    source_sha = _write_source(source)

    with pytest.raises(module.BuildError, match="SOURCE_IDENTITY_HOLD"):
        module.build_draft(source, tmp_path / "wrong-sha.xlsx", expected_source_sha256="0" * 64)
    with pytest.raises(module.BuildError, match="OUTPUT_SOURCE_COLLISION"):
        module.build_draft(source, source, expected_source_sha256=source_sha)
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"existing")
    with pytest.raises(module.BuildError, match="OUTPUT_EXISTS"):
        module.build_draft(source, existing, expected_source_sha256=source_sha)

    duplicate_source = tmp_path / "duplicate.xlsx"
    _write_source(duplicate_source, identities=IDENTITIES[:-1] + (IDENTITIES[-2],))
    duplicate_sha = hashlib.sha256(duplicate_source.read_bytes()).hexdigest()
    with pytest.raises(module.BuildError, match="SOURCE_IDENTITY_HOLD"):
        module.build_draft(duplicate_source, tmp_path / "duplicate-output.xlsx", expected_source_sha256=duplicate_sha)

    short_source = tmp_path / "short.xlsx"
    short_sha = _write_source(short_source, identities=IDENTITIES[:-1])
    with pytest.raises(module.BuildError, match="SOURCE_IDENTITY_HOLD"):
        module.build_draft(short_source, tmp_path / "short-output.xlsx", expected_source_sha256=short_sha)


def test_builder_writes_formula_like_source_values_as_literal_text(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bản tin điện tử"
    sheet.append(list(HEADERS))
    values = [None] * len(HEADERS)
    values[0] = "1. Thông báo mời thầu"
    values[1] = None
    values[4] = "Gói (Số thông báo: IB2600000001-00)"
    values[7] = "1.000.000"
    sheet.append(values)
    sheet["B2"].value = "=SUM(1,1)"
    sheet["B2"].data_type = "s"
    workbook.save(source)
    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    output = tmp_path / "draft.xlsx"

    module.build_draft(source, output, expected_source_sha256=source_sha, expected_source_rows=1, expected_identity_count=1)

    check = load_workbook(output, data_only=False)
    cell = check["01_NGUON_11_GOI"]["B2"]
    assert cell.value == "=SUM(1,1)"
    assert cell.data_type == "s"
    check.close()


def test_builder_validation_failure_leaves_no_partial_final_output(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    output = tmp_path / "draft.xlsx"
    source_sha = _write_source(source)

    def fail_final_validation(path, **kwargs):
        if Path(path).resolve() == output.resolve():
            raise module.BuildError("ARTIFACT_VALIDATION_HOLD: synthetic final failure")

    monkeypatch.setattr(module, "validate_output_workbook", fail_final_validation)
    with pytest.raises(module.BuildError, match="ARTIFACT_VALIDATION_HOLD"):
        module.build_draft(source, output, expected_source_sha256=source_sha)
    assert not output.exists()
    assert not list(tmp_path.glob(".draft.xlsx.*"))


def test_profile_v1_formalizes_roles_and_keeps_ground_truth_blank(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    old_draft = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT.xlsx"
    profile_output = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT_PROFILE_V1.xlsx"
    source_sha = _write_source(source)

    module.build_draft(source, old_draft, expected_source_sha256=source_sha)
    old_sha = hashlib.sha256(old_draft.read_bytes()).hexdigest()
    result = module.build_profile_v1(source, profile_output, expected_source_sha256=source_sha)

    assert result.output == profile_output.resolve()
    workbook = load_workbook(profile_output, data_only=False)
    profile = workbook["00_HUONG_DAN_PROFILE"]
    assert profile["B2"].value == "GOLDEN_CANDIDATE_01"
    assert profile["B3"].value == "DRAFT"
    assert profile["B8"].value == "APPROVED"
    assert profile["B9"].value == "NOT_APPROVED"
    assert profile["B10"].value == "QI_TBMT_OPPORTUNITY_SCREENING"
    assert profile["B11"].value == "1.0"
    assert "Lọc và xếp thứ tự các TBMT" in profile["B16"].value
    profile_text = " ".join(str(cell.value or "") for row in profile.iter_rows() for cell in row)
    assert "SOLUTION APPROVAL" in profile_text

    catalog_headers = {cell.value: cell.column for cell in profile[43]}
    roles = {
        profile.cell(row, catalog_headers["CRITERION_ID"]).value: profile.cell(
            row, catalog_headers["QI_ROLE_DECISION"]
        ).value
        for row in range(44, 55)
    }
    assert roles == {
        "C01": "INACTIVE",
        "C02": "INACTIVE",
        "C03": "INACTIVE",
        "C04": "INACTIVE",
        "C05": "INACTIVE",
        "C06": "INACTIVE",
        "C07": "HARD",
        "C08": "INACTIVE",
        "C09": "PREFERENCE",
        "C10": "PREFERENCE",
        "C11": "INACTIVE",
    }
    allowed = {
        profile.cell(row, catalog_headers["CRITERION_ID"]).value: profile.cell(
            row, catalog_headers["QI_ALLOWED_VALUE_OR_THRESHOLD"]
        ).value
        for row in range(44, 55)
    }
    assert "2,000,000,000 VND" in allowed["C09"]
    assert ">2B" in allowed["C09"]
    actual_regions = tuple(part.strip() for part in allowed["C10"].split(";"))
    assert actual_regions == EXPECTED_HUMAN_REGIONS
    assert not any(region in allowed["C10"] for region in UNAUTHORIZED_REGIONS)
    assert "100 km" not in profile_text

    labels = workbook["02_TEAM_BID_GAN_NHAN"]
    label_headers = {cell.value: cell.column for cell in labels[1]}
    assert all(labels.cell(row, label_headers["EXPECTED_SCREENING"]).value in (None, "") for row in range(2, 13))
    assert all(labels.cell(row, label_headers["EXPECTED_PRIORITY"]).value in (None, "") for row in range(2, 13))
    priority_formulas = [
        str(validation.formula1)
        for validation in labels.data_validations.dataValidation
        if validation.type == "list"
    ]
    assert any("HIGH,MEDIUM,UNRANKED" in formula for formula in priority_formulas)
    assert all("A,B" not in formula for formula in priority_formulas)

    evidence = workbook["03_BANG_CHUNG_TIEU_CHI"]
    evidence_headers = {cell.value: cell.column for cell in evidence[1]}
    evidence_roles = [
        evidence.cell(row, evidence_headers["PROFILE_ROLE"]).value for row in range(2, 123)
    ]
    assert evidence_roles.count("HARD") == 11
    assert evidence_roles.count("PREFERENCE") == 22
    assert evidence_roles.count("INACTIVE") == 88
    assert all(evidence.cell(row, evidence_headers["EXPECTED_OUTCOME"]).value in (None, "") for row in range(2, 123))
    assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    workbook.close()

    assert hashlib.sha256(old_draft.read_bytes()).hexdigest() == old_sha


def test_profile_v1_region_correction_creates_separate_marked_artifact_and_preserves_bad_profile(
    tmp_path: Path,
) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    old_draft = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT.xlsx"
    bad_profile = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT_PROFILE_V1.xlsx"
    corrected_profile = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT_PROFILE_V1_CORR1.xlsx"
    source_sha = _write_source(source)

    module.build_draft(source, old_draft, expected_source_sha256=source_sha)
    module.build_profile_v1(source, bad_profile, expected_source_sha256=source_sha)
    bad_profile_sha = hashlib.sha256(bad_profile.read_bytes()).hexdigest()
    result = module.build_profile_v1_correction(source, corrected_profile, expected_source_sha256=source_sha)

    assert result.output == corrected_profile.resolve()
    assert hashlib.sha256(bad_profile.read_bytes()).hexdigest() == bad_profile_sha
    workbook = load_workbook(corrected_profile, data_only=False)
    profile = workbook["00_HUONG_DAN_PROFILE"]
    profile_text = " ".join(str(cell.value or "") for row in profile.iter_rows() for cell in row)
    assert "CORR-GT01-PROFILE-V1-REGION-01" in profile_text
    assert "100 km" not in profile_text
    workbook.close()


def test_profile_v1_rejects_existing_target_without_mutating_old_draft(tmp_path: Path) -> None:
    module = _builder_module()
    source = tmp_path / "source.xlsx"
    old_draft = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT.xlsx"
    target = tmp_path / "GOLDEN_CANDIDATE_01_DRAFT_PROFILE_V1.xlsx"
    source_sha = _write_source(source)
    module.build_draft(source, old_draft, expected_source_sha256=source_sha)
    old_sha = hashlib.sha256(old_draft.read_bytes()).hexdigest()
    target.write_bytes(b"existing-profile")

    with pytest.raises(module.BuildError, match="OUTPUT_EXISTS"):
        module.build_profile_v1(source, target, expected_source_sha256=source_sha)

    assert target.read_bytes() == b"existing-profile"
    assert hashlib.sha256(old_draft.read_bytes()).hexdigest() == old_sha
