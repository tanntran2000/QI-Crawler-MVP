"""Build the Team Bid labeling workbook for Golden Candidate 01.

This is preparation tooling only.  It copies source facts and provenance into
an auditable workbook while leaving all business decisions empty for Team Bid.
"""

import argparse
import hashlib
import os
import sys
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[2]
SRC_ROOT = REPO_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from qi_crawler.market_intelligence.opportunity_contract import OpportunitySourceType
from qi_crawler.market_intelligence.opportunity_intelligence import (
    OpportunityIntelligenceService,
)

EXPECTED_SOURCE_SHA256 = "3fa652eba17710647f29537d89a1b2c6e0f16660c7d0b0a5cccec05657b206d0"
EXPECTED_SOURCE_ROWS = 11
EXPECTED_IDENTITY_COUNT = 11
CRITERION_COUNT = 11
SOURCE_HEADER_SCAN_LIMIT = 50
SHEET_NAMES = (
    "00_HUONG_DAN_PROFILE",
    "01_NGUON_11_GOI",
    "02_TEAM_BID_GAN_NHAN",
    "03_BANG_CHUNG_TIEU_CHI",
)
SOURCE_REQUIRED_HEADERS = {
    "GÓI THẦU",
    "HÌNH THỨC LỰA CHỌN NHÀ THẦU",
}
SOURCE_FORMULA_PREFIXES = ("=", "+", "-", "@")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="000000")
THIN_GREY = Side(style="thin", color="D9E1F2")


class BuildError(RuntimeError):
    """A bounded, user-readable failure from the draft builder."""

    def __init__(self, message: str) -> None:
        self.code = message.split(":", maxsplit=1)[0]
        super().__init__(message)


@dataclass(frozen=True, slots=True)
class SourceLayout:
    sheet: str
    header_row: int
    header_columns: dict[str, int]


@dataclass(frozen=True, slots=True)
class SourceBundle:
    path: Path
    sha256: str
    size: int
    source_sheet: str
    source_header_row: int
    source_headers: tuple[str, ...]
    source_data_rows: int
    successfully_imported_rows: int
    rejected_or_skipped_rows: int
    import_issue_count: int
    exact_identity_count: int
    items: tuple[Any, ...]
    layout: SourceLayout


@dataclass(frozen=True, slots=True)
class DraftBuildResult:
    output: Path
    source_sha_before: str
    source_sha_after: str
    source_size_before: int
    source_size_after: int
    source_sheet: str
    source_data_rows: int
    successfully_imported_rows: int
    rejected_or_skipped_rows: int
    import_issue_count: int
    exact_identity_count: int
    output_sha256: str
    output_size: int


@dataclass(frozen=True, slots=True)
class CriterionSpec:
    criterion_id: str
    name: str
    source_availability: str
    source_fields: str
    current_evaluator: str
    implementation_phase: str
    reviewer_suggestion: str
    missing_data_policy: str
    notes: str


CRITERIA = (
    CriterionSpec(
        "C01",
        "NOTICE_STATUS",
        "MISSING",
        "",
        "NO",
        "PARK",
        "Status active/open candidate",
        "UNKNOWN",
        "No authoritative notice-status field in this source.",
    ),
    CriterionSpec(
        "C02",
        "DEADLINE_REMAINING",
        "DIRECT",
        "THỜI GIAN ĐÓNG THẦU(HẠN CUỐI TIẾP NHẬN BG)",
        "NO",
        "S1_CANDIDATE",
        "Deadline and preparation-time candidate",
        "UNKNOWN",
        "Raw deadline is preserved; evaluation_at and day policy require Human input.",
    ),
    CriterionSpec(
        "C03",
        "E_BIDDING_MODE",
        "MISSING",
        "",
        "NO",
        "S1_CANDIDATE",
        "Qua mạng / Không qua mạng candidate",
        "UNKNOWN",
        "Do not infer from the publication URL or issuing address.",
    ),
    CriterionSpec(
        "C04",
        "SELECTION_METHOD",
        "DIRECT",
        "HÌNH THỨC LỰA CHỌN NHÀ THẦU",
        "YES",
        "CURRENT",
        "ĐTRR / CHCT preference candidate",
        "UNKNOWN",
        "Normalized value is copied only as a technical source projection.",
    ),
    CriterionSpec(
        "C05",
        "PROCUREMENT_METHOD",
        "DIRECT",
        "PHƯƠNG THỨC LỰA CHỌN NHÀ THẦU",
        "NO",
        "S1_CANDIDATE",
        "1G1B / other procurement-method candidate",
        "UNKNOWN",
        "Kept separate from selection_method; no filter is implemented here.",
    ),
    CriterionSpec(
        "C06",
        "PROCUREMENT_DOMAIN",
        "CONTEXTUAL",
        "DỰ ÁN; GÓI THẦU; NỘI DUNG CHÍNH CỦA GÓI THẦU",
        "NO",
        "S1_CANDIDATE",
        "Goods / non-consulting / mixed candidate",
        "UNKNOWN",
        "Only raw business text is provided; no domain label is assigned.",
    ),
    CriterionSpec(
        "C07",
        "TOPIC",
        "CONTEXTUAL",
        "DỰ ÁN; GÓI THẦU; NỘI DUNG CHÍNH CỦA GÓI THẦU",
        "NO",
        "S2_CANDIDATE",
        "IT / Network / Electronics / Software candidate",
        "UNKNOWN",
        "Raw text is evidence only; no topic classification is assigned.",
    ),
    CriterionSpec(
        "C08",
        "ACTIVITY",
        "CONTEXTUAL",
        "DỰ ÁN; GÓI THẦU; NỘI DUNG CHÍNH CỦA GÓI THẦU",
        "NO",
        "S2_CANDIDATE",
        "Supply / installation / training / other candidate",
        "UNKNOWN",
        "Raw text is evidence only; no activity classification is assigned.",
    ),
    CriterionSpec(
        "C09",
        "BUDGET",
        "DIRECT",
        "GIÁ GÓI THẦU",
        "YES",
        "CURRENT",
        "Budget threshold candidate",
        "UNKNOWN",
        "No budget threshold is invented in this draft.",
    ),
    CriterionSpec(
        "C10",
        "EXECUTION_LOCATION",
        "MISSING",
        "",
        "YES",
        "CURRENT",
        "Execution-location candidate",
        "UNKNOWN",
        "Execution location is not inferred from the procuring-entity address.",
    ),
    CriterionSpec(
        "C11",
        "BUYER_SEGMENT",
        "CONTEXTUAL",
        "BÊN MỜI THẦU",
        "NO",
        "S2_CANDIDATE",
        "Buyer-segment preference candidate",
        "UNKNOWN",
        "Organization name is preserved without risk or reputation inference.",
    ),
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_cell(cell: Any, value: Any) -> None:
    """Write typed values while forcing source-like strings to literal text."""

    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _compact(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return " ".join(text.split()) or None


def _find_source_layout(path: Path, sheet_name: str) -> SourceLayout:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise BuildError("SOURCE_IDENTITY_HOLD: imported sheet is missing")
        sheet = workbook[sheet_name]
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=SOURCE_HEADER_SCAN_LIMIT, values_only=True),
            start=1,
        ):
            header_columns = {
                str(value).strip(): index
                for index, value in enumerate(row, start=1)
                if value not in (None, "")
            }
            if SOURCE_REQUIRED_HEADERS.issubset(header_columns):
                return SourceLayout(sheet_name, row_number, header_columns)
    finally:
        workbook.close()
    raise BuildError("SOURCE_IDENTITY_HOLD: source header row is missing")


def _load_source(
    source_path: str | Path,
    *,
    expected_source_sha256: str = EXPECTED_SOURCE_SHA256,
    expected_source_rows: int = EXPECTED_SOURCE_ROWS,
    expected_identity_count: int = EXPECTED_IDENTITY_COUNT,
) -> SourceBundle:
    path = Path(source_path).resolve()
    if path.suffix.lower() != ".xlsx" or not path.is_file():
        raise BuildError("SOURCE_IDENTITY_HOLD: source must be an existing .xlsx file")
    source_sha = _sha256(path)
    source_size = path.stat().st_size
    if source_sha.casefold() != expected_source_sha256.casefold():
        raise BuildError(
            f"SOURCE_IDENTITY_HOLD: expected SHA {expected_source_sha256}, got {source_sha}"
        )

    service = OpportunityIntelligenceService(None)
    try:
        loaded = service.load_workbook(path, OpportunitySourceType.TBMT)
    except Exception as exc:
        raise BuildError(f"SOURCE_IDENTITY_HOLD: TBMT import failed: {type(exc).__name__}") from exc

    if loaded.source_row_count != expected_source_rows:
        raise BuildError(
            f"SOURCE_IDENTITY_HOLD: expected {expected_source_rows} data rows, got {loaded.source_row_count}"
        )
    if len(loaded.items) != expected_source_rows:
        raise BuildError(
            f"SOURCE_IDENTITY_HOLD: expected {expected_source_rows} imported candidates, got {len(loaded.items)}"
        )
    if loaded.issues:
        raise BuildError(
            f"SOURCE_IDENTITY_HOLD: import produced {len(loaded.issues)} issue(s)"
        )

    identities = [(item.identity.base_id, item.identity.revision) for item in loaded.items]
    if len(set(identities)) != expected_identity_count:
        raise BuildError(
            f"SOURCE_IDENTITY_HOLD: expected {expected_identity_count} unique exact identities, got {len(set(identities))}"
        )
    for item in loaded.items:
        if item.source_type is not OpportunitySourceType.TBMT:
            raise BuildError("SOURCE_IDENTITY_HOLD: imported source type is not TBMT")
        if item.source_sha256.casefold() != source_sha.casefold():
            raise BuildError("SOURCE_IDENTITY_HOLD: item source SHA differs from source")
        if item.provenance.get("source_sha256", "").casefold() != source_sha.casefold():
            raise BuildError("SOURCE_IDENTITY_HOLD: provenance SHA differs from source")
        if item.provenance.get("sheet") != loaded.sheet:
            raise BuildError("SOURCE_IDENTITY_HOLD: provenance sheet differs from source")
        if item.provenance.get("source_row") != item.source_row:
            raise BuildError("SOURCE_IDENTITY_HOLD: provenance row differs from item row")

    layout = _find_source_layout(path, loaded.sheet)
    return SourceBundle(
        path=path,
        sha256=source_sha,
        size=source_size,
        source_sheet=loaded.sheet,
        source_header_row=layout.header_row,
        source_headers=tuple(loaded.headers),
        source_data_rows=loaded.source_row_count,
        successfully_imported_rows=len(loaded.items),
        rejected_or_skipped_rows=loaded.source_row_count - len(loaded.items),
        import_issue_count=len(loaded.issues),
        exact_identity_count=len(set(identities)),
        items=tuple(loaded.items),
        layout=layout,
    )


def _style_table(sheet: Any, header_row: int = 1) -> None:
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GREY)
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_cells in sheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max((len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells), default=0)
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)


def _add_list_validation(sheet: Any, column: int, first_row: int, last_row: int, values: Iterable[str]) -> None:
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"{get_column_letter(column)}{first_row}:{get_column_letter(column)}{last_row}")


def _profile_sheet(workbook: Workbook, source: SourceBundle) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[0])
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "GOLDEN_CANDIDATE_01 — DRAFT"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E78")
    header_values = (
        ("CANDIDATE_ID", "GOLDEN_CANDIDATE_01"),
        ("STATUS", "DRAFT"),
        ("SOURCE_FILENAME", source.path.name),
        ("SOURCE_SHA256", source.sha256),
        ("SOURCE_TYPE", "TBMT"),
        ("SOURCE_ROW_COUNT", source.source_data_rows),
        ("PROFILE_STATUS", "NOT_APPROVED"),
        ("GROUND_TRUTH_STATUS", "NOT_APPROVED"),
        ("PROFILE_ID", ""),
        ("PROFILE_VERSION", ""),
        ("TIMEZONE", "Asia/Ho_Chi_Minh"),
        ("EVALUATION_AT", ""),
        ("DEADLINE_DAY_MODE", ""),
        ("DEADLINE_BOUNDARY_POLICY", ""),
        ("BUSINESS_OBJECTIVE", ""),
        ("AUTHORIZED_LABELER", ""),
        ("APPROVER", ""),
    )
    for row, (label, value) in enumerate(header_values, start=2):
        _write_cell(sheet.cell(row, 1), label)
        _write_cell(sheet.cell(row, 2), value)
        sheet.cell(row, 1).font = Font(name="Arial", size=10, bold=True, color="1F4E78")
        if label in {"PROFILE_ID", "PROFILE_VERSION", "EVALUATION_AT", "DEADLINE_DAY_MODE", "DEADLINE_BOUNDARY_POLICY", "BUSINESS_OBJECTIVE", "AUTHORIZED_LABELER", "APPROVER"}:
            sheet.cell(row, 2).fill = INPUT_FILL

    _write_cell(sheet["A20"], "BLIND LABELING INSTRUCTION")
    sheet["A20"].font = Font(name="Arial", size=10, bold=True, color="1F4E78")
    _write_cell(
        sheet["B20"],
        "Team Bid should label the candidate before viewing Crawler screening predictions for the same profile.",
    )
    sheet["B20"].alignment = Alignment(wrap_text=True, vertical="top")
    _write_cell(sheet["A22"], "AGGREGATION RULE — DOCUMENT ONLY")
    sheet["A22"].font = Font(name="Arial", size=10, bold=True, color="1F4E78")
    for row, rule in enumerate(
        (
            "NO ACTIVE HARD CRITERIA → UNFILTERED",
            "ANY HARD FAIL → EXCLUDE",
            "NO HARD FAIL + ANY HARD UNKNOWN → NEEDS_REVIEW",
            "ALL ACTIVE HARD PASS → SELECT",
        ),
        start=23,
    ):
        _write_cell(sheet.cell(row, 2), rule)

    catalog_headers = (
        "CRITERION_ID",
        "CRITERION_NAME",
        "SOURCE_AVAILABILITY",
        "SOURCE_FIELD_OR_FIELDS",
        "CURRENT_EVALUATOR",
        "IMPLEMENTATION_PHASE",
        "REVIEWER_SUGGESTION",
        "QI_ROLE_DECISION",
        "QI_ALLOWED_VALUE_OR_THRESHOLD",
        "QI_PRIORITY_EFFECT",
        "MISSING_DATA_POLICY",
        "QI_NOTES",
    )
    catalog_row = 29
    for column, value in enumerate(catalog_headers, start=1):
        _write_cell(sheet.cell(catalog_row, column), value)
    for row, criterion in enumerate(CRITERIA, start=catalog_row + 1):
        values = (
            criterion.criterion_id,
            criterion.name,
            criterion.source_availability,
            criterion.source_fields,
            criterion.current_evaluator,
            criterion.implementation_phase,
            criterion.reviewer_suggestion,
            "",
            "",
            "",
            criterion.missing_data_policy,
            criterion.notes,
        )
        for column, value in enumerate(values, start=1):
            _write_cell(sheet.cell(row, column), value)
    _add_list_validation(sheet, 8, catalog_row + 1, catalog_row + len(CRITERIA), ("HARD", "PREFERENCE", "INACTIVE"))
    _style_table(sheet, catalog_row)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 72
    for row in (20, 22):
        sheet.row_dimensions[row].height = 30


def _cell_coordinate(layout: SourceLayout, field: str, source_row: int) -> str:
    column = layout.header_columns.get(field)
    return f"{get_column_letter(column)}{source_row}" if column else ""


def _business_text(item: Any) -> tuple[str, str, str]:
    fields = item.raw_fields
    names = ("DỰ ÁN", "GÓI THẦU", "NỘI DUNG CHÍNH CỦA GÓI THẦU")
    values = tuple(_compact(fields.get(name)) or "" for name in names)
    excerpt = " | ".join(value for value in values if value)
    return "; ".join(name for name, value in zip(names, values) if value), "; ".join(
        _cell_name for _cell_name, value in zip(names, values) if value
    ), excerpt


def _evidence_for(item: Any, criterion: CriterionSpec, layout: SourceLayout) -> tuple[str, str, str, str, str]:
    fields = item.raw_fields
    source_row = item.source_row
    direct_fields = {
        "C02": "THỜI GIAN ĐÓNG THẦU(HẠN CUỐI TIẾP NHẬN BG)",
        "C04": "HÌNH THỨC LỰA CHỌN NHÀ THẦU",
        "C05": "PHƯƠNG THỨC LỰA CHỌN NHÀ THẦU",
        "C09": "GIÁ GÓI THẦU",
        "C11": "BÊN MỜI THẦU",
    }
    if criterion.criterion_id in direct_fields:
        field = direct_fields[criterion.criterion_id]
        value = _compact(fields.get(field)) or ""
        authority = "SOURCE_DIRECT" if value else "MISSING_SOURCE"
        return value, field, _cell_coordinate(layout, field, source_row), "", authority
    if criterion.criterion_id in {"C06", "C07", "C08"}:
        source_fields, _, excerpt = _business_text(item)
        coordinates = "; ".join(
            _cell_coordinate(layout, field, source_row)
            for field in ("DỰ ÁN", "GÓI THẦU", "NỘI DUNG CHÍNH CỦA GÓI THẦU")
            if field in layout.header_columns and _compact(fields.get(field))
        )
        return excerpt, source_fields, coordinates, "", "SOURCE_CONTEXTUAL" if excerpt else "MISSING_SOURCE"
    return "", "", "", "", "MISSING_SOURCE"


def _source_sheet(workbook: Workbook, source: SourceBundle) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[1])
    headers = source.source_headers + (
        "QI_IB_RAW",
        "QI_BASE_ID",
        "QI_REVISION",
        "QI_SOURCE_SHEET",
        "QI_SOURCE_ROW",
        "QI_SOURCE_SHA256",
        "QI_OBSERVATION_KEY",
        "QI_SELECTION_METHOD_NORMALIZED",
    )
    for column, value in enumerate(headers, start=1):
        _write_cell(sheet.cell(1, column), value)
    for row, item in enumerate(source.items, start=2):
        values = [item.raw_fields.get(header) for header in source.source_headers]
        values.extend(
            (
                item.identity.raw_id,
                item.identity.base_id,
                item.identity.revision,
                item.sheet,
                item.source_row,
                item.source_sha256,
                item.observation_key,
                item.selection_method,
            )
        )
        for column, value in enumerate(values, start=1):
            _write_cell(sheet.cell(row, column), value)
        for column in range(len(source.source_headers) + 1, len(headers) + 1):
            sheet.cell(row, column).number_format = "@"
    _style_table(sheet)


def _label_sheet(workbook: Workbook, source: SourceBundle) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[2])
    headers = (
        "IB",
        "BASE_ID",
        "REVISION",
        "SOURCE_ROW",
        "OBSERVATION_KEY",
        "EXPECTED_SCREENING",
        "EXPECTED_PRIORITY",
        "BUSINESS_REASON",
        "MISSING_INFORMATION",
        "LABELER",
        "LABELLED_AT",
        "APPROVER",
        "APPROVED_AT",
        "LABEL_STATUS",
        "LABEL_VERSION",
    )
    for column, value in enumerate(headers, start=1):
        _write_cell(sheet.cell(1, column), value)
    for row, item in enumerate(source.items, start=2):
        values = (
            item.identity.raw_id,
            item.identity.base_id,
            item.identity.revision,
            item.source_row,
            item.observation_key,
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "DRAFT",
            "",
        )
        for column, value in enumerate(values, start=1):
            _write_cell(sheet.cell(row, column), value)
    _add_list_validation(sheet, 6, 2, source.source_data_rows + 1, ("SELECT", "NEEDS_REVIEW", "EXCLUDE", "UNFILTERED"))
    _add_list_validation(sheet, 7, 2, source.source_data_rows + 1, ("A", "B", "UNRANKED"))
    _add_list_validation(sheet, 14, 2, source.source_data_rows + 1, ("DRAFT", "DISPUTED", "APPROVED"))
    _style_table(sheet)


def _evidence_sheet(workbook: Workbook, source: SourceBundle) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[3])
    headers = (
        "IB",
        "REVISION",
        "SOURCE_ROW",
        "OBSERVATION_KEY",
        "CRITERION_ID",
        "CRITERION_NAME",
        "PROFILE_ROLE",
        "SOURCE_AVAILABILITY",
        "OBSERVED_VALUE",
        "EXPECTED_OUTCOME",
        "SOURCE_FIELD",
        "SOURCE_CELL_OR_COORDINATE",
        "ADDITIONAL_SOURCE_REFS",
        "SOURCE_EXCERPT",
        "EVIDENCE_AUTHORITY",
        "EXCLUSION_BASIS",
        "TEAM_BID_EXPLANATION",
        "LABEL_STATUS",
    )
    for column, value in enumerate(headers, start=1):
        _write_cell(sheet.cell(1, column), value)
    row = 2
    for item in source.items:
        for criterion in CRITERIA:
            observed, source_field, coordinate, additional_refs, authority = _evidence_for(item, criterion, source.layout)
            values = (
                item.identity.raw_id,
                item.identity.revision,
                item.source_row,
                item.observation_key,
                criterion.criterion_id,
                criterion.name,
                "PENDING_PROFILE",
                criterion.source_availability,
                observed,
                "",
                source_field,
                coordinate,
                additional_refs,
                observed,
                authority,
                "",
                "",
                "DRAFT",
            )
            for column, value in enumerate(values, start=1):
                _write_cell(sheet.cell(row, column), value)
            row += 1
    _add_list_validation(sheet, 10, 2, row - 1, ("PASS", "FAIL", "UNKNOWN"))
    _add_list_validation(sheet, 15, 2, row - 1, ("SOURCE_DIRECT", "SOURCE_CONTEXTUAL", "MISSING_SOURCE", "HUMAN_BUSINESS_JUDGMENT"))
    _add_list_validation(sheet, 18, 2, row - 1, ("DRAFT", "DISPUTED", "APPROVED"))
    _style_table(sheet)


def _build_workbook(source: SourceBundle) -> Workbook:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    _profile_sheet(workbook, source)
    _source_sheet(workbook, source)
    _label_sheet(workbook, source)
    _evidence_sheet(workbook, source)
    return workbook


def validate_output_workbook(
    output_path: str | Path,
    *,
    expected_sha256: str,
    expected_source_rows: int,
    expected_identity_count: int,
) -> None:
    path = Path(output_path).resolve()
    if not path.is_file():
        raise BuildError("ARTIFACT_VALIDATION_HOLD: output workbook is missing")
    workbook = load_workbook(path, data_only=False)
    try:
        if workbook.sheetnames != list(SHEET_NAMES):
            raise BuildError("ARTIFACT_VALIDATION_HOLD: sheet order is incorrect")
        profile = workbook[SHEET_NAMES[0]]
        if profile["B2"].value != "GOLDEN_CANDIDATE_01" or profile["B3"].value != "DRAFT":
            raise BuildError("ARTIFACT_VALIDATION_HOLD: profile header is incorrect")
        if profile["B5"].value.casefold() != expected_sha256.casefold():
            raise BuildError("ARTIFACT_VALIDATION_HOLD: profile SHA is incorrect")
        if profile["B8"].value != "NOT_APPROVED" or profile["B9"].value != "NOT_APPROVED":
            raise BuildError("ARTIFACT_VALIDATION_HOLD: approval state is not draft")

        source_sheet = workbook[SHEET_NAMES[1]]
        source_headers = [cell.value for cell in source_sheet[1]]
        metadata_start = source_headers.index("QI_IB_RAW")
        if source_sheet.max_row - 1 != expected_source_rows:
            raise BuildError("ARTIFACT_VALIDATION_HOLD: source row count is incorrect")
        source_rows = list(source_sheet.iter_rows(min_row=2, values_only=False))
        identity_column = source_headers.index("QI_IB_RAW")
        identities = {row[identity_column].value for row in source_rows}
        if len(identities) != expected_identity_count:
            raise BuildError("ARTIFACT_VALIDATION_HOLD: source identity count is incorrect")
        row_column = source_headers.index("QI_SOURCE_ROW")
        sha_column = source_headers.index("QI_SOURCE_SHA256")
        for row in source_rows:
            if row[sha_column].value.casefold() != expected_sha256.casefold():
                raise BuildError("ARTIFACT_VALIDATION_HOLD: source row SHA is incorrect")
            if row[row_column].value is None:
                raise BuildError("ARTIFACT_VALIDATION_HOLD: source row coordinate is missing")
        if metadata_start != len(source_headers) - 8:
            raise BuildError("ARTIFACT_VALIDATION_HOLD: source metadata position is incorrect")

        labels = workbook[SHEET_NAMES[2]]
        if labels.max_row - 1 != expected_source_rows:
            raise BuildError("ARTIFACT_VALIDATION_HOLD: label row count is incorrect")
        label_headers = {cell.value: cell.column for cell in labels[1]}
        for row in range(2, labels.max_row + 1):
            if labels.cell(row, label_headers["EXPECTED_SCREENING"]).value not in (None, ""):
                raise BuildError("ARTIFACT_VALIDATION_HOLD: screening prediction leaked")
            if labels.cell(row, label_headers["EXPECTED_PRIORITY"]).value not in (None, ""):
                raise BuildError("ARTIFACT_VALIDATION_HOLD: priority prediction leaked")
            if labels.cell(row, label_headers["LABEL_STATUS"]).value != "DRAFT":
                raise BuildError("ARTIFACT_VALIDATION_HOLD: label lifecycle state is incorrect")

        evidence = workbook[SHEET_NAMES[3]]
        if evidence.max_row - 1 != expected_source_rows * CRITERION_COUNT:
            raise BuildError("ARTIFACT_VALIDATION_HOLD: criterion row count is incorrect")
        evidence_headers = {cell.value: cell.column for cell in evidence[1]}
        for row in range(2, evidence.max_row + 1):
            if evidence.cell(row, evidence_headers["EXPECTED_OUTCOME"]).value not in (None, ""):
                raise BuildError("ARTIFACT_VALIDATION_HOLD: criterion prediction leaked")
            if evidence.cell(row, evidence_headers["PROFILE_ROLE"]).value != "PENDING_PROFILE":
                raise BuildError("ARTIFACT_VALIDATION_HOLD: profile role was prefilled")

        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise BuildError("ARTIFACT_VALIDATION_HOLD: unexpected formula cell")
    finally:
        workbook.close()


def build_draft(
    source_path: str | Path,
    output_path: str | Path,
    *,
    expected_source_sha256: str = EXPECTED_SOURCE_SHA256,
    expected_source_rows: int = EXPECTED_SOURCE_ROWS,
    expected_identity_count: int = EXPECTED_IDENTITY_COUNT,
) -> DraftBuildResult:
    source = Path(source_path).resolve()
    output = Path(output_path).resolve()
    if source == output:
        raise BuildError("OUTPUT_SOURCE_COLLISION: output path equals source path")
    if output.exists():
        raise BuildError("OUTPUT_EXISTS: output already exists; choose another destination")

    bundle = _load_source(
        source,
        expected_source_sha256=expected_source_sha256,
        expected_source_rows=expected_source_rows,
        expected_identity_count=expected_identity_count,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output.name}.",
            suffix=".xlsx",
            dir=output.parent,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook = _build_workbook(bundle)
        try:
            workbook.save(temporary_path)
        finally:
            workbook.close()

        source_sha_before_replace = _sha256(source)
        source_size_before_replace = source.stat().st_size
        if (
            source_sha_before_replace.casefold() != bundle.sha256.casefold()
            or source_size_before_replace != bundle.size
        ):
            raise BuildError("SOURCE_MUTATION_HOLD: source changed during build")

        validate_output_workbook(
            temporary_path,
            expected_sha256=bundle.sha256,
            expected_source_rows=bundle.source_data_rows,
            expected_identity_count=bundle.exact_identity_count,
        )
        os.replace(temporary_path, output)
        temporary_path = None

        source_sha_after = _sha256(source)
        source_size_after = source.stat().st_size
        if source_sha_after.casefold() != bundle.sha256.casefold() or source_size_after != bundle.size:
            output.unlink(missing_ok=True)
            raise BuildError("SOURCE_MUTATION_HOLD: source changed after build")
        validate_output_workbook(
            output,
            expected_sha256=bundle.sha256,
            expected_source_rows=bundle.source_data_rows,
            expected_identity_count=bundle.exact_identity_count,
        )
        return DraftBuildResult(
            output=output,
            source_sha_before=bundle.sha256,
            source_sha_after=source_sha_after,
            source_size_before=bundle.size,
            source_size_after=source_size_after,
            source_sheet=bundle.source_sheet,
            source_data_rows=bundle.source_data_rows,
            successfully_imported_rows=bundle.successfully_imported_rows,
            rejected_or_skipped_rows=bundle.rejected_or_skipped_rows,
            import_issue_count=bundle.import_issue_count,
            exact_identity_count=bundle.exact_identity_count,
            output_sha256=_sha256(output),
            output_size=output.stat().st_size,
        )
    except BuildError:
        if output.exists() and temporary_path is not None:
            output.unlink(missing_ok=True)
        raise
    except Exception as exc:
        if output.exists() and temporary_path is not None:
            output.unlink(missing_ok=True)
        raise BuildError(f"ARTIFACT_VALIDATION_HOLD: {type(exc).__name__}") from exc
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Build GOLDEN_CANDIDATE_01 — DRAFT")
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    try:
        result = build_draft(args.source, args.output)
    except BuildError as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print(f"DRAFT_BUILT={result.output}")
    print(f"SOURCE_SHA256={result.source_sha_after}")
    print(f"SOURCE_ROWS={result.source_data_rows}")
    print(f"IDENTITIES={result.exact_identity_count}")
    print(f"OUTPUT_SHA256={result.output_sha256}")
    print(f"OUTPUT_SIZE={result.output_size}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
