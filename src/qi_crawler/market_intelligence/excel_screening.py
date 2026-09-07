"""Deterministic, source-preserving Excel screening for Team Bid.

This module is deliberately additive.  It does not change the generic MI
filter engine or turn a screening result into a human decision/Ground Truth.
Every non-empty row after the selected header is retained as either an
accounted data row or an explicitly classified non-record row.
"""

from __future__ import annotations

import hashlib
import re
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal
from enum import StrEnum
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from qi_crawler.excel_safety import safe_excel_row

from .khmt_normalization import parse_package_price, parse_plan_identity
from .tbmt_normalization import parse_tbmt_notice_identity

PROFILE_VERSION = "QI_TBMT_OPPORTUNITY_SCREENING/1.0"
C07_RULESET_VERSION = "C07-V1-R07-2026-09"
LOCATION_HINT_VERSION = "LOCATION-HINT-V1-2026-09"
MAX_HEADER_SCAN_ROWS = 100
PRICE_LIMIT = Decimal(2000000000)


class C07Result(StrEnum):
    PASS = "PASS"
    UNKNOWN = "UNKNOWN"
    FAIL = "FAIL"


class C09Result(StrEnum):
    PASS = "PASS"
    UNKNOWN = "UNKNOWN"
    FAIL = "FAIL"


class C10Result(StrEnum):
    PASS = "PASS"
    UNKNOWN = "UNKNOWN"
    FAIL = "FAIL"


class ScreeningResult(StrEnum):
    SELECT = "SELECT"
    NEEDS_REVIEW = "NEEDS_REVIEW"
    EXCLUDE = "EXCLUDE"


class Priority(StrEnum):
    HIGH = "HIGH"
    MEDIUM = "MEDIUM"
    UNRANKED = "UNRANKED"


class RecordStatus(StrEnum):
    READ_OK = "READ_OK"
    READ_ERROR = "READ_ERROR"


class SourceType(StrEnum):
    KHMT = "KHMT"
    TBMT = "TBMT"
    UNKNOWN = "UNKNOWN"


@dataclass(frozen=True, slots=True)
class SourceDetection:
    source_type: SourceType
    filename_hint: SourceType
    identity_namespace: str | None
    identity_values: tuple[str, ...]
    detected_header_row: int
    sheet: str
    filename_schema_conflict: bool
    evidence: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class NonRecordRow:
    source_row: int
    reason_code: str
    raw_values: tuple[Any, ...]


@dataclass(frozen=True, slots=True)
class ScreeningRecord:
    source_row: int
    status: RecordStatus
    identity: str | None
    package_name: str | None
    package_price_raw: str | None
    package_price: Decimal | None
    c07: C07Result | None
    c09: C09Result | None
    c10: C10Result | None
    screening: ScreeningResult | None
    priority: Priority | None
    c07_rule_code: str | None
    evidence_field: str | None
    evidence_excerpt: str | None
    location_hint: str | None
    location_hint_basis: str | None
    location_hint_field: str | None
    location_hint_evidence: str | None
    source_type: SourceType
    source_filename: str
    source_sha256: str
    source_sheet: str
    source_row_error_code: str | None = None
    source_row_error: str | None = None
    missing_information: str | None = None
    raw_fields: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class ScreeningRun:
    source_path: Path
    source_filename: str
    source_sha256: str
    source_type: SourceType
    identity_namespace: str | None
    filename_hint: SourceType
    filename_schema_conflict: bool
    sheet: str
    detected_header_row: int
    records: tuple[ScreeningRecord, ...]
    non_record_rows_detail: tuple[NonRecordRow, ...]
    post_header_nonempty_rows: int
    data_record_rows: int
    read_ok_rows: int
    read_error_rows: int

    @property
    def non_record_rows(self) -> int:
        return len(self.non_record_rows_detail)

    @property
    def select_count(self) -> int:
        return sum(record.screening is ScreeningResult.SELECT for record in self.records)

    @property
    def needs_review_count(self) -> int:
        return sum(record.screening is ScreeningResult.NEEDS_REVIEW for record in self.records)

    @property
    def exclude_count(self) -> int:
        return sum(record.screening is ScreeningResult.EXCLUDE for record in self.records)

    @property
    def read_error_count(self) -> int:
        return self.read_error_rows

    @property
    def accounting_invariant_status(self) -> str:
        if self.data_record_rows != self.read_ok_rows + self.read_error_rows:
            return "FAIL"
        if self.post_header_nonempty_rows != self.data_record_rows + self.non_record_rows:
            return "FAIL"
        return "PASS"


_IDENTITY_RE = re.compile(
    r"\b(?P<namespace>PL|IB)\s*(?P<number>\d{8,14})\s*-\s*(?P<revision>[0-9A-Za-z]{2})\b",
    re.IGNORECASE,
)

_LOCATION_ALIASES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("TPHCM", ("tphcm", "tp hcm", "tp.hcm", "tp. hcm", "thành phố hồ chí minh", "ho chi minh city")),
    ("Đồng Nai", ("đồng nai", "dong nai")),
    ("Tây Ninh", ("tây ninh", "tay ninh")),
    ("Bình Dương", ("bình dương", "binh duong")),
    ("Long An", ("long an",)),
    ("Bà Rịa - Vũng Tàu", ("bà rịa", "ba ria", "vũng tàu", "vung tau")),
    ("Bình Phước", ("bình phước", "binh phuoc")),
    ("Lâm Đồng", ("lâm đồng", "lam dong")),
    ("Phan Thiết", ("phan thiết", "phan thiet")),
    ("Cần Giờ", ("cần giờ", "can gio")),
    ("Hà Nội", ("hà nội", "ha noi")),
)

_C07_RULES: tuple[tuple[str, str, tuple[str, ...]], ...] = (
    (
        "R07-COMPUTE",
        "COMPUTE",
        (
            r"\bm[aá]y tính\b",
            r"\blaptop\b",
            r"\bworkstation\b",
            r"\bm[aá]y chủ\b",
            r"\bserver\b",
            r"\bthiết bị lưu trữ\b",
            r"\bm[aá]y in\b",
            r"\bm[aá]y qu[eé]t\b",
        ),
    ),
    (
        "R07-NETWORK",
        "NETWORK",
        (
            r"m[aạ]ng\s+lan",
            r"wi[- ]?fi",
            r"hạ tầng mạng",
            r"thiết bị mạng",
            r"\bswitch\b",
            r"\brouter\b",
            r"access\s+point",
            r"hệ thống mạng",
            r"cáp mạng",
        ),
    ),
    (
        "R07-SOFTWARE-LICENSE",
        "SOFTWARE-LICENSE",
        (
            r"phần mềm",
            r"bản quyền phần mềm",
            r"\blicense\b",
            r"\bwindows\b",
            r"nâng cấp phần mềm",
            r"bảo trì phần mềm",
            r"triển khai phần mềm",
        ),
    ),
    (
        "R07-CYBERSECURITY",
        "CYBERSECURITY",
        (
            r"an toàn thông tin",
            r"\battt\b",
            r"an ninh mạng",
            r"\bfirewall\b",
            r"\bsoc\b",
            r"\bsiem\b",
            r"giải pháp bảo mật",
            r"cybersecurity",
        ),
    ),
    (
        "R07-SURVEILLANCE",
        "SURVEILLANCE",
        (r"\bcamera\b", r"giám sát hình ảnh bằng camera"),
    ),
    (
        "R07-DIGITAL-SYSTEM",
        "DIGITAL-SYSTEM",
        (
            r"\bscada\b",
            r"hệ thống thông tin",
            r"hệ thống báo cáo điện tử",
            r"hệ thống cntt",
            r"giải pháp số",
        ),
    ),
    (
        "R07-IT-SERVICE",
        "IT-SERVICE",
        (
            r"thuê dịch vụ cntt",
            r"cung cấp dịch vụ cntt",
            r"triển khai giải pháp chuyển đổi số",
            r"thuê giải pháp chuyển đổi số",
        ),
    ),
)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).strip().split())


def _key(value: Any) -> str:
    return _text(value).upper()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _filename_hint(path: Path) -> SourceType:
    stem = path.stem.upper()
    if stem.startswith("KHMT"):
        return SourceType.KHMT
    if stem.startswith("TBMT"):
        return SourceType.TBMT
    return SourceType.UNKNOWN


def _identity_values(values: Iterable[Any]) -> tuple[tuple[str, ...], str | None]:
    found: list[str] = []
    namespaces: set[str] = set()
    for value in values:
        for match in _IDENTITY_RE.finditer(_text(value)):
            namespace = match.group("namespace").upper()
            canonical = f"{namespace}{match.group('number')}-{match.group('revision')}"
            namespaces.add(namespace)
            if canonical not in found:
                found.append(canonical)
    namespace = next(iter(namespaces)) if len(namespaces) == 1 else ("MIXED" if namespaces else None)
    return tuple(found), namespace


def _header_type(headers: set[str]) -> SourceType:
    khmt = {"SỐ KẾ HOẠCH", "TÊN GÓI THẦU"}.issubset(headers)
    tbmt = {"BÊN MỜI THẦU", "DỰ ÁN", "GÓI THẦU", "GIÁ GÓI THẦU"}.issubset(headers)
    if khmt and not tbmt:
        return SourceType.KHMT
    if tbmt and not khmt:
        return SourceType.TBMT
    return SourceType.UNKNOWN


def _detect_header(workbook: Any) -> tuple[Any, int, tuple[str | None, ...], SourceType]:
    best: tuple[int, int, Any, tuple[str | None, ...], SourceType] | None = None
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=MAX_HEADER_SCAN_ROWS, values_only=True),
            start=1,
        ):
            headers = {_key(value) for value in row if _key(value)}
            kind = _header_type(headers)
            score = len(headers.intersection({"SỐ KẾ HOẠCH", "TÊN GÓI THẦU", "BÊN MỜI THẦU", "DỰ ÁN", "GÓI THẦU", "GIÁ GÓI THẦU"}))
            candidate = (1 if kind is not SourceType.UNKNOWN else 0, score, sheet, tuple(_text(value) or None for value in row), kind)
            if best is None or candidate[:2] > best[:2]:
                best = candidate
            if kind is not SourceType.UNKNOWN:
                return sheet, row_number, candidate[3], kind
    if best is None or best[4] is SourceType.UNKNOWN:
        raise ValueError("No supported KHMT/TBMT table header was found")
    return best[2], 0, best[3], best[4]


def _canonical_fields(headers: tuple[str | None, ...], values: tuple[Any, ...]) -> dict[str, Any]:
    fields: dict[str, Any] = {}
    for index, header in enumerate(headers):
        if header is None:
            continue
        fields[_key(header)] = values[index] if index < len(values) else None
    return fields


def _first(fields: dict[str, Any], *names: str) -> Any:
    for name in names:
        value = fields.get(name)
        if value not in (None, ""):
            return value
    return None


def _is_non_record(values: tuple[Any, ...]) -> str | None:
    texts = [_text(value) for value in values if _text(value)]
    joined = " ".join(texts).upper()
    if not joined:
        return "EMPTY_ROW"
    if re.search(r"\b(TỔNG CỘNG|TỔNG SỐ|TỔNG CỘNG CÁC GÓI|GRAND TOTAL)\b", joined):
        return "EXPLICIT_TOTAL_FOOTER"
    if re.match(r"^(GHI CHÚ|NOTE|NOTES|BÁO CÁO|REPORT)\s*[:：]", joined):
        return "EXPLICIT_REPORT_NOTE"
    return None


def _evidence_excerpt(text: str, match: re.Match[str]) -> str:
    start = max(0, match.start() - 60)
    end = min(len(text), match.end() + 60)
    return text[start:end].strip()


def _screen_c07(source_type: SourceType, fields: dict[str, Any]) -> tuple[C07Result, str | None, str | None, str | None, str | None]:
    names = ("TÊN GÓI THẦU", "GÓI THẦU")
    evidence_fields = names if source_type is SourceType.KHMT else (*names, "NỘI DUNG CHÍNH CỦA GÓI THẦU")
    for field_name in evidence_fields:
        value = _text(fields.get(field_name))
        if not value:
            continue
        for rule_code, _label, patterns in _C07_RULES:
            for pattern in patterns:
                match = re.search(pattern, value, flags=re.IGNORECASE)
                if match:
                    return C07Result.PASS, rule_code, field_name, _evidence_excerpt(value, match), None
    return C07Result.UNKNOWN, None, None, None, "Thiếu bằng chứng công nghệ ở trường package-specific được tin cậy"


def _location_hint(fields: dict[str, Any], source_type: SourceType) -> tuple[str | None, str | None, str | None, str | None]:
    field_names = (
        ("TÊN CHỦ ĐẦU TƯ", "TÊN DỰ ÁN", "TÊN GÓI THẦU", "NỘI DUNG PHÊ DUYỆT")
        if source_type is SourceType.KHMT
        else ("BÊN MỜI THẦU", "DỰ ÁN", "GÓI THẦU", "ĐỊA CHỈ BÊN MỜI THẦU")
    )
    matches: list[tuple[str, str, str]] = []
    for field_name in field_names:
        value = _text(fields.get(field_name))
        if not value:
            continue
        folded = value.casefold()
        for label, aliases in _LOCATION_ALIASES:
            if any(alias.casefold() in folded for alias in aliases):
                matches.append((label, field_name, value))
    if not matches:
        return None, None, None, None
    labels: list[str] = []
    for label, _field_name, _value in matches:
        if label not in labels:
            labels.append(label)
    basis = "; ".join(f"{field_name}: {value}" for _label, field_name, value in matches)
    first = matches[0]
    return ", ".join(labels), "SOURCE_TEXT_LOCATION_HINT", first[1], basis


def _screen_c09(value: Any) -> tuple[C09Result, Decimal | None, str | None]:
    raw = _text(value) or None
    price = parse_package_price(value)
    if price is None:
        return C09Result.UNKNOWN, None, raw
    return (C09Result.PASS if price <= PRICE_LIMIT else C09Result.FAIL), price, raw


def _record_from_row(
    *,
    source_type: SourceType,
    fields: dict[str, Any],
    source_row: int,
    source_filename: str,
    source_sha256: str,
    sheet: str,
) -> ScreeningRecord:
    identity_raw = _first(fields, "SỐ KẾ HOẠCH", "GÓI THẦU")
    package_raw = _first(fields, "TÊN GÓI THẦU", "GÓI THẦU")
    package_name = _text(package_raw) or None
    row_identity_text = _text(identity_raw)
    row_identity_namespace = (
        _IDENTITY_RE.search(row_identity_text).group("namespace").upper()
        if _IDENTITY_RE.search(row_identity_text)
        else None
    )
    if source_type is SourceType.KHMT or row_identity_namespace == "PL":
        identity_obj = parse_plan_identity(identity_raw)
    elif source_type is SourceType.TBMT or row_identity_namespace == "IB":
        identity_obj = parse_tbmt_notice_identity(identity_raw)
    else:
        identity_obj = None
    if identity_obj is None:
        identity = _text(identity_raw) or None
    else:
        identity = identity_obj.raw if source_type is SourceType.KHMT else identity_obj.raw_id
    errors: list[tuple[str, str]] = []
    if identity_obj is None:
        errors.append(("INVALID_IDENTITY", "Identity is missing or malformed for detected source namespace"))
    if package_name is None:
        errors.append(("EMPTY_PACKAGE_NAME", "Package name is missing"))
    if errors:
        return ScreeningRecord(
            source_row=source_row,
            status=RecordStatus.READ_ERROR,
            identity=identity,
            package_name=package_name,
            package_price_raw=_text(_first(fields, "GIÁ GÓI THẦU")) or None,
            package_price=None,
            c07=None,
            c09=None,
            c10=None,
            screening=None,
            priority=None,
            c07_rule_code=None,
            evidence_field=None,
            evidence_excerpt=None,
            location_hint=None,
            location_hint_basis=None,
            location_hint_field=None,
            location_hint_evidence=None,
            source_type=source_type,
            source_filename=source_filename,
            source_sha256=source_sha256,
            source_sheet=sheet,
            source_row_error_code=errors[0][0],
            source_row_error="; ".join(message for _code, message in errors),
            missing_information="; ".join(code for code, _message in errors),
            raw_fields=fields,
        )

    c07, rule_code, evidence_field, evidence_excerpt, missing = _screen_c07(source_type, fields)
    c09, price, price_raw = _screen_c09(_first(fields, "GIÁ GÓI THẦU"))
    hints = _location_hint(fields, source_type)
    screening = ScreeningResult.SELECT if c07 is C07Result.PASS else ScreeningResult.NEEDS_REVIEW
    priority = Priority.MEDIUM if screening is ScreeningResult.SELECT else Priority.UNRANKED
    return ScreeningRecord(
        source_row=source_row,
        status=RecordStatus.READ_OK,
        identity=identity,
        package_name=package_name,
        package_price_raw=price_raw,
        package_price=price,
        c07=c07,
        c09=c09,
        c10=C10Result.UNKNOWN,
        screening=screening,
        priority=priority,
        c07_rule_code=rule_code,
        evidence_field=evidence_field,
        evidence_excerpt=evidence_excerpt,
        location_hint=hints[0],
        location_hint_basis=hints[1],
        location_hint_field=hints[2],
        location_hint_evidence=hints[3],
        source_type=source_type,
        source_filename=source_filename,
        source_sha256=source_sha256,
        source_sheet=sheet,
        missing_information=missing,
        raw_fields=fields,
    )


def screen_excel_workbook(path: str | Path) -> ScreeningRun:
    """Read one workbook into a complete, provenance-preserving screening run."""

    source_path = Path(path).resolve()
    if source_path.suffix.casefold() != ".xlsx" or not source_path.is_file():
        raise ValueError("Excel screening requires an existing .xlsx workbook")
    source_sha256 = _sha256(source_path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        sheet, header_row, headers, schema_type = _detect_header(workbook)
        rows_for_detection: list[Any] = []
        for row in sheet.iter_rows(values_only=True):
            rows_for_detection.extend(value for value in row if value not in (None, ""))
        _identities, identity_namespace = _identity_values(rows_for_detection)
        identity_type = (
            SourceType.KHMT
            if identity_namespace == "PL"
            else SourceType.TBMT
            if identity_namespace == "IB"
            else SourceType.UNKNOWN
        )
        source_type = (
            SourceType.UNKNOWN
            if identity_namespace == "MIXED"
            else identity_type
            if identity_type is not SourceType.UNKNOWN
            else schema_type
        )
        filename_hint = _filename_hint(source_path)
        filename_schema_conflict = filename_hint is not SourceType.UNKNOWN and source_type is not SourceType.UNKNOWN and filename_hint is not source_type
        records: list[ScreeningRecord] = []
        non_records: list[NonRecordRow] = []
        post_header_nonempty_rows = 0
        for source_row, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            values_tuple = tuple(values)
            if not any(value not in (None, "") for value in values_tuple):
                continue
            post_header_nonempty_rows += 1
            non_record_reason = _is_non_record(values_tuple)
            if non_record_reason is not None and non_record_reason != "EMPTY_ROW":
                non_records.append(NonRecordRow(source_row, non_record_reason, values_tuple))
                continue
            fields = _canonical_fields(headers, values_tuple)
            records.append(
                _record_from_row(
                    source_type=source_type,
                    fields=fields,
                    source_row=source_row,
                    source_filename=source_path.name,
                    source_sha256=source_sha256,
                    sheet=sheet.title,
                )
            )
        return ScreeningRun(
            source_path=source_path,
            source_filename=source_path.name,
            source_sha256=source_sha256,
            source_type=source_type,
            identity_namespace=identity_namespace,
            filename_hint=filename_hint,
            filename_schema_conflict=filename_schema_conflict,
            sheet=sheet.title,
            detected_header_row=header_row,
            records=tuple(records),
            non_record_rows_detail=tuple(non_records),
            post_header_nonempty_rows=post_header_nonempty_rows,
            data_record_rows=len(records),
            read_ok_rows=sum(record.status is RecordStatus.READ_OK for record in records),
            read_error_rows=sum(record.status is RecordStatus.READ_ERROR for record in records),
        )
    finally:
        workbook.close()


_BUSINESS_HEADERS = (
    "Tên gói",
    "Giá gói",
    "Lý do phù hợp",
    "Địa bàn gợi ý",
    "Cần kiểm tra gì",
    "Mã PL/IB",
    "C07",
    "C09",
    "C10",
    "Screening Result",
    "Priority",
    "C07 Rule Code",
    "Evidence Field",
    "Evidence Excerpt",
    "Location Hint Basis",
    "Source Type",
    "Revision",
    "Source Sheet",
    "Source Row",
    "Source SHA256 / bounded trace reference",
)

_AUDIT_HEADERS = (
    "Accounting Disposition",
    "Source Row",
    "Mã PL/IB",
    "Tên gói",
    "Giá gói",
    "C07",
    "C09",
    "C10",
    "Screening Result",
    "Priority",
    "Error Code",
    "Error",
    "Missing Information",
    "Source Type",
    "Source Sheet",
    "Source SHA256",
)


def _revision(identity: str | None) -> str | None:
    return identity.rsplit("-", 1)[1] if identity and "-" in identity else None


def _reason(record: ScreeningRecord) -> str:
    if record.status is RecordStatus.READ_ERROR:
        return record.source_row_error or "Read error"
    if record.c07 is C07Result.PASS:
        return record.evidence_excerpt or "Package-specific technology evidence"
    return "No bounded package-specific technology evidence"


def _check_before_export(run: ScreeningRun) -> None:
    if run.accounting_invariant_status != "PASS":
        raise ValueError("EXPORT_ACCEPTANCE=FAIL: source accounting invariant is not satisfied")


def _record_values(record: ScreeningRecord) -> tuple[Any, ...]:
    return (
        record.package_name,
        record.package_price,
        _reason(record),
        record.location_hint,
        record.missing_information,
        record.identity,
        record.c07.value if record.c07 else None,
        record.c09.value if record.c09 else None,
        record.c10.value if record.c10 else None,
        record.screening.value if record.screening else None,
        record.priority.value if record.priority else None,
        record.c07_rule_code,
        record.evidence_field,
        record.evidence_excerpt,
        record.location_hint_basis,
        record.source_type.value,
        _revision(record.identity),
        record.source_sheet,
        record.source_row,
        record.source_sha256,
    )


def _ordered_select(records: Iterable[ScreeningRecord]) -> list[ScreeningRecord]:
    c09_rank = {C09Result.PASS: 0, C09Result.UNKNOWN: 1, C09Result.FAIL: 2, None: 3}
    return sorted(
        (record for record in records if record.status is RecordStatus.READ_OK and record.screening is ScreeningResult.SELECT),
        key=lambda record: (c09_rank[record.c09], record.source_row),
    )


def _ordered_review(records: Iterable[ScreeningRecord]) -> list[ScreeningRecord]:
    return [
        record
        for record in records
        if record.status is RecordStatus.READ_OK and record.screening is ScreeningResult.NEEDS_REVIEW
    ]


def export_screening_workbook(run: ScreeningRun, output_path: str | Path) -> Path:
    """Export exactly the four Team Bid screening sheets."""

    _check_before_export(run)
    destination = Path(output_path).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    first = workbook.active
    first.title = "01_CO_HOI_XEM_XET"
    for sheet_name in ("02_CAN_BO_SUNG", "03_AUDIT_TOAN_BO", "04_META"):
        workbook.create_sheet(sheet_name)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    def write_business_sheet(sheet: Any, rows: Iterable[ScreeningRecord]) -> None:
        sheet.append(_BUSINESS_HEADERS)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for record in rows:
            sheet.append(safe_excel_row(_record_values(record)))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    write_business_sheet(first, _ordered_select(run.records))
    write_business_sheet(workbook["02_CAN_BO_SUNG"], _ordered_review(run.records))

    audit = workbook["03_AUDIT_TOAN_BO"]
    audit.append(_AUDIT_HEADERS)
    for cell in audit[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for record in run.records:
        audit.append(
            safe_excel_row(
                (
                record.status.value,
                record.source_row,
                record.identity,
                record.package_name,
                record.package_price,
                record.c07.value if record.c07 else None,
                record.c09.value if record.c09 else None,
                record.c10.value if record.c10 else None,
                record.screening.value if record.screening else None,
                record.priority.value if record.priority else None,
                record.source_row_error_code,
                record.source_row_error,
                record.missing_information,
                record.source_type.value,
                record.source_sheet,
                record.source_sha256,
                )
            )
        )
    for non_record in run.non_record_rows_detail:
        audit.append(
            safe_excel_row(
                (
                "NON_RECORD",
                non_record.source_row,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                None,
                non_record.reason_code,
                "Row retained outside DATA_RECORD_ROWS",
                None,
                run.source_type.value,
                run.sheet,
                run.source_sha256,
                )
            )
        )
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions

    meta = workbook["04_META"]
    meta.append(("FIELD", "VALUE"))
    for cell in meta[1]:
        cell.fill = header_fill
        cell.font = header_font
    metadata = (
        ("SOURCE_FILENAME", run.source_filename),
        ("SOURCE_SHA256", run.source_sha256),
        ("SHEET", run.sheet),
        ("DETECTED_HEADER_ROW", run.detected_header_row),
        ("DETECTED_SOURCE_TYPE", run.source_type.value),
        ("IDENTITY_NAMESPACE", run.identity_namespace),
        ("FILENAME_HINT", run.filename_hint.value),
        ("FILENAME_SCHEMA_CONFLICT", run.filename_schema_conflict),
        ("POST_HEADER_NONEMPTY_ROWS", run.post_header_nonempty_rows),
        ("DATA_RECORD_ROWS", run.data_record_rows),
        ("NON_RECORD_ROWS", run.non_record_rows),
        ("READ_OK_ROWS", run.read_ok_rows),
        ("READ_ERROR_ROWS", run.read_error_rows),
        ("SELECT_COUNT", run.select_count),
        ("NEEDS_REVIEW_COUNT", run.needs_review_count),
        ("EXCLUDE_COUNT", run.exclude_count),
        ("READ_ERROR_COUNT", run.read_error_count),
        ("PROFILE_VERSION", PROFILE_VERSION),
        ("C07_RULESET_VERSION", C07_RULESET_VERSION),
        ("LOCATION_HINT_VERSION", LOCATION_HINT_VERSION),
        ("ACCOUNTING_INVARIANT_STATUS", run.accounting_invariant_status),
        ("EXPORT_ACCEPTANCE", "PASS" if run.accounting_invariant_status == "PASS" else "FAIL"),
        ("MACHINE_OUTPUT_GROUND_TRUTH", "NO"),
    )
    for row in metadata:
        meta.append(safe_excel_row(row))
    meta.freeze_panes = "A2"
    meta.auto_filter.ref = meta.dimensions

    for sheet in workbook.worksheets:
        for column_cells in sheet.columns:
            width = min(60, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in sheet.iter_rows():
            for cell in row:
                if cell.column in (1, 3, 4, 5, 13, 14, 15):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(destination)
    workbook.close()
    return destination


def screen_and_export(path: str | Path, output_path: str | Path) -> ScreeningRun:
    run = screen_excel_workbook(path)
    export_screening_workbook(run, output_path)
    return run
